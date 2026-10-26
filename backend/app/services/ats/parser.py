import io
import re
from typing import Any

import docx
import pdfplumber

from app.core.config import settings
from app.core.errors import ValidationError

# Magic byte signatures
PDF_MAGIC = b"%PDF"
DOCX_MAGIC = b"PK\x03\x04"

SECTION_PATTERNS = {
    "education": re.compile(r"^\s*(education|academic\s+background|academics|qualifications)\s*$", re.IGNORECASE),
    "experience": re.compile(r"^\s*(experience|work\s+experience|professional\s+experience|employment\s+history|internships)\s*$", re.IGNORECASE),
    "projects": re.compile(r"^\s*(projects|technical\s+projects|academic\s+projects|key\s+projects)\s*$", re.IGNORECASE),
    "skills": re.compile(r"^\s*(skills|technical\s+skills|core\s+competencies|technologies|tools\s+&\s+technologies)\s*$", re.IGNORECASE),
    "certifications": re.compile(r"^\s*(certifications|licenses|courses|achievements|honors\s+&\s+awards)\s*$", re.IGNORECASE),
}

EMAIL_REGEX = re.compile(r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+")
PHONE_REGEX = re.compile(r"(?:\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}")
URL_REGEX = re.compile(r"https?://(?:www\.)?[-a-zA-Z0-9@:%._\+~#=]{1,256}\.[a-zA-Z0-9()]{1,6}\b(?:[-a-zA-Z0-9()@:%_\+.~#?&//=]*)")

def validate_magic_bytes_and_size(data: bytes, filename: str) -> str:
    max_bytes = settings.MAX_UPLOAD_MB * 1024 * 1024
    if len(data) > max_bytes:
        raise ValidationError(f"File size exceeds maximum allowed size of {settings.MAX_UPLOAD_MB} MB")
    
    if len(data) == 0:
        raise ValidationError("Uploaded file is empty")

    if data.startswith(PDF_MAGIC):
        return "pdf"
    elif data.startswith(DOCX_MAGIC):
        return "docx"
    
    # Try decoding as plaintext utf-8
    try:
        data.decode("utf-8")
        return "text"
    except UnicodeDecodeError:
        raise ValidationError("Unsupported file format. Please upload a valid PDF, DOCX, or text file.")

def extract_text_from_pdf(data: bytes) -> str:
    text_parts = []
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
    except Exception as e:
        raise ValidationError(f"Failed to extract text from PDF: {e!s}")
    
    if not text_parts:
        raise ValidationError("No readable text found in PDF document")
    return "\n".join(text_parts)

def extract_text_from_docx(data: bytes) -> str:
    text_parts = []
    try:
        doc = docx.Document(io.BytesIO(data))
        for para in doc.paragraphs:
            if para.text:
                text_parts.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    text_parts.append(" | ".join(row_text))
    except Exception as e:
        raise ValidationError(f"Failed to extract text from DOCX: {e!s}")
    
    if not text_parts:
        raise ValidationError("No readable text found in DOCX document")
    return "\n".join(text_parts)

def normalize_text(raw_text: str) -> str:
    # Fix hyphenated words across lines (e.g. multi-\nthreading -> multithreading)
    text = re.sub(r"(\w+)-\s*\n\s*(\w+)", r"\1\2", raw_text)
    # Normalize multiple newlines
    text = re.sub(r"\n{3,}", "\n\n", text)
    # Normalize spaces and tabs
    text = re.sub(r"[ \t]+", " ", text)
    return text.strip()

def parse_sections_and_contact(normalized_text: str) -> dict[str, Any]:
    lines = normalized_text.splitlines()
    sections: dict[str, list[str]] = {
        "header": [],
        "education": [],
        "experience": [],
        "projects": [],
        "skills": [],
        "certifications": [],
        "other": [],
    }

    current_section = "header"
    for line in lines:
        clean_line = line.strip()
        if not clean_line:
            continue

        matched_section = None
        for sec_name, pattern in SECTION_PATTERNS.items():
            if pattern.match(clean_line):
                matched_section = sec_name
                break

        if matched_section:
            current_section = matched_section
        else:
            sections[current_section].append(clean_line)

    # Flatten section contents into text
    parsed_sections = {k: "\n".join(v) for k, v in sections.items() if v}

    # Extract contact info
    emails = EMAIL_REGEX.findall(normalized_text)
    phones = PHONE_REGEX.findall(normalized_text)
    urls = URL_REGEX.findall(normalized_text)

    contact_info = {
        "email": emails[0] if emails else None,
        "phone": phones[0] if phones else None,
        "urls": urls,
        "has_contact": bool(emails or phones),
    }

    # Extract bullets
    bullet_lines = []
    for line in lines:
        stripped = line.strip()
        # Common bullet markers: -, *, •, ⁃, 1., etc.
        if re.match(r"^([•\-\*⁃–—]|\d+\.)\s+", stripped):
            clean_bullet = re.sub(r"^([•\-\*⁃–—]|\d+\.)\s+", "", stripped)
            if len(clean_bullet) > 15:
                bullet_lines.append(clean_bullet)

    return {
        "sections": parsed_sections,
        "contact_info": contact_info,
        "bullets": bullet_lines,
        "word_count": len(normalized_text.split()),
    }

def parse_resume_content(raw_bytes: bytes, filename: str) -> tuple[str, str, dict[str, Any]]:
    file_type = validate_magic_bytes_and_size(raw_bytes, filename)
    if file_type == "pdf":
        raw_text = extract_text_from_pdf(raw_bytes)
    elif file_type == "docx":
        raw_text = extract_text_from_docx(raw_bytes)
    else:
        raw_text = raw_bytes.decode("utf-8", errors="replace")

    normalized = normalize_text(raw_text)
    parsed_data = parse_sections_and_contact(normalized)
    return file_type, normalized, parsed_data
